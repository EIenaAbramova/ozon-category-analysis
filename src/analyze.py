"""Ozon category analysis. Python 3.11+; run from any directory."""
from pathlib import Path
import argparse, hashlib, json, gzip
import pandas as pd
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
COLS = ['cat1','cat2','cat3','cat4','name','url','super','seller','brand','scheme',
        'revenue','price','availability','daily_revenue','daily_units','lost_revenue',
        'warehouses','delivery','catalog_views','card_views','catalog_conversion','card_conversion','drr','created']

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--source',type=Path)
    args=ap.parse_args()
    (ROOT/'outputs').mkdir(exist_ok=True); (ROOT/'data').mkdir(exist_ok=True)
    cache=ROOT/'data'/'products.csv.gz'
    if args.source:
        source=args.source
        wb=openpyxl.load_workbook(source,read_only=True,data_only=True)
        ws=wb['ТОП товаров']
        df=pd.DataFrame(ws.iter_rows(min_row=10,max_col=24,values_only=True),columns=COLS)
        df.insert(0,'source_row',np.arange(10,10+len(df)))
        df.to_csv(cache,index=False,compression='gzip',date_format='%Y-%m-%d')
        meta={'source':source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
              'sheet':'ТОП товаров','header_row':9,'period':'03.07.2024–30.07.2024',
              'column_names':dict(zip(COLS,next(ws.iter_rows(min_row=9,max_row=9,values_only=True)))),
              'analysis_date':'2026-10-17'}
        (ROOT/'data'/'source.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
        wb.close()
    else:
        files=[cache] if cache.exists() else sorted((ROOT/'data').glob('products_*.csv.gz'))
        if not files: raise FileNotFoundError('No data CSV files; supply --source original.xlsx')
        df=pd.concat([pd.read_csv(f,keep_default_na=False,na_values=['']) for f in files],ignore_index=True)
    for c in COLS[10:23]: df[c]=pd.to_numeric(df[c],errors='coerce')
    empty=df[COLS].isna().all(axis=1)
    invalid=df.revenue.isna() | (df.revenue<0) | df.drr.isna() | (df.drr<0) | df.cat4.isna()
    # Missingness and anomalous values are retained in diagnostics; not replaced by zero.
    valid=df.loc[~empty & ~invalid].copy()
    dup=valid.duplicated(subset=['url'],keep=False)
    diagnostics={'rows':len(df),'empty_rows':int(empty.sum()),'invalid_rows':int(invalid.sum()),
                 'duplicate_url_rows':int(dup.sum()),'unique_urls':int(valid.url.nunique()),
                 'exact_duplicate_rows':int(valid[COLS].duplicated().sum()),
                 'missing':{c:int(df[c].isna().sum()) for c in COLS},
                 'drr_over_1':int((df.drr>1).sum()),'drr_zero':int((valid.drr==0).sum()),
                 'revenue_zero':int((valid.revenue==0).sum()),
                 'min_revenue':float(valid.revenue.min()),'max_revenue':float(valid.revenue.max())}
    if dup.any():
        raise ValueError('Duplicate product URLs: inspect grain before aggregation')
    valid['million']=valid.revenue>1_000_000
    valid['ad_spend_est']=valid.revenue*valid.drr
    keys=['cat1','cat2','cat3','cat4']
    g=valid.groupby(keys,dropna=False)
    cats=g.agg(n=('url','size'),million_n=('million','sum'),revenue=('revenue','sum'),
               ad_spend_est=('ad_spend_est','sum'),mean_drr=('drr','mean'),median_drr=('drr','median'),
               zero_drr_n=('drr',lambda x:(x==0).sum()),sellers=('seller','nunique'),
               median_price=('price','median'),median_availability=('availability','median'),
               lost_revenue=('lost_revenue','sum')).reset_index()
    cats['million_share']=cats.million_n/cats.n
    cats['weighted_drr']=cats.ad_spend_est/cats.revenue
    top_seller=valid.groupby(keys+['seller'],dropna=False).revenue.sum().groupby(keys).max()
    cats['top_seller_share']=cats.set_index(keys).index.map(top_seller)/cats.revenue
    baseline={'n':len(valid),'revenue':float(valid.revenue.sum()),'million_n':int(valid.million.sum()),
              'million_share':float(valid.million.mean()),'weighted_drr':float(valid.ad_spend_est.sum()/valid.revenue.sum()),
              'categories':len(cats)}
    eligible=cats.n>=30
    # Explicit analytical cutoffs; not requirements supplied by the employer.
    # A quartile share threshold is unsuitable here: its value is zero.
    drr_cut=.03
    share_cut=.10
    cats['selected']=eligible & (cats.weighted_drr<=drr_cut) & (cats.million_share>=share_cut)
    # Pareto frontier: no category with both lower DRR and higher share.
    universe=cats.loc[eligible]
    cats['pareto']=False
    for idx,r in universe.iterrows():
        dominates=(universe.weighted_drr<=r.weighted_drr)&(universe.million_share>=r.million_share)&((universe.weighted_drr<r.weighted_drr)|(universe.million_share>r.million_share))
        cats.loc[idx,'pareto']=not dominates.any()
    cats=cats.sort_values(['selected','million_share','weighted_drr','n'],ascending=[False,False,True,False])
    cats.insert(0,'category_id',cats[keys].astype(str).agg(' / '.join,axis=1))
    cats.to_csv(ROOT/'outputs'/'categories.csv',index=False,encoding='utf-8-sig')
    broad=valid.groupby('cat1').agg(n=('url','size'),million_n=('million','sum'),revenue=('revenue','sum'),ad_spend_est=('ad_spend_est','sum')).reset_index()
    broad['million_share']=broad.million_n/broad.n; broad['weighted_drr']=broad.ad_spend_est/broad.revenue
    broad.to_csv(ROOT/'outputs'/'categories_level1.csv',index=False,encoding='utf-8-sig')
    selected=cats[cats.selected]
    examples=valid.merge(selected[keys],on=keys).query('revenue > 1000000 and drr <= @drr_cut')
    examples=examples.sort_values(['drr','revenue'],ascending=[True,False]).groupby(keys,dropna=False).head(3)
    examples.to_csv(ROOT/'outputs'/'product_examples.csv',index=False,encoding='utf-8-sig')
    sensitivity=[]
    for min_n in [10,30,100]:
        for dc,sc,rule in [(.01,.10,'ДРР ≤1%, доля ≥10%'),(.03,.10,'ДРР ≤3%, доля ≥10%'),(.03,.05,'ДРР ≤3%, доля ≥5%')]:
            q=cats[(cats.n>=min_n)&(cats.weighted_drr<=dc)&(cats.million_share>=sc)]
            sensitivity.append({'min_n':min_n,'drr_max':dc,'share_min':sc,'rule':rule,'categories':len(q),'products':int(q.n.sum())})
    summary={'baseline':baseline,'diagnostics':diagnostics,'thresholds':{'min_n':30,'drr_max':drr_cut,'million_share_min':share_cut,'eligible_categories':int(eligible.sum())},
             'selected':selected.replace({np.nan:None}).to_dict('records'),
             'lowest_drr':cats.sort_values(['weighted_drr','million_share'],ascending=[True,False]).head(20).replace({np.nan:None}).to_dict('records'),
             'highest_share':cats.sort_values(['million_share','n'],ascending=False).head(20).replace({np.nan:None}).to_dict('records'),
             'sensitivity':sensitivity}
    (ROOT/'outputs'/'analysis.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    # Independent checks: additive controls and direct row filtering for every selected category.
    assert cats.n.sum()==len(valid)
    assert cats.million_n.sum()==valid.million.sum()
    assert np.isclose(cats.revenue.sum(),valid.revenue.sum())
    for _,r in selected.iterrows():
        rows=g.get_group(tuple(r[k] for k in keys))
        assert len(rows)==r.n and int((rows.revenue>1_000_000).sum())==r.million_n
        assert np.isclose(np.average(rows.drr,weights=rows.revenue),r.weighted_drr)
    print(json.dumps({'baseline':baseline,'diagnostics':diagnostics,'thresholds':summary['thresholds'],'selected':selected[['cat4','n','million_n','million_share','weighted_drr','top_seller_share']].to_dict('records'),'checks':'passed'},ensure_ascii=False,indent=2))

if __name__=='__main__': main()
