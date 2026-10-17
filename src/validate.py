"""Independent controls using standard-library CSV streaming plus XLSX inspection."""
from pathlib import Path
from collections import defaultdict
import csv,gzip,json,math
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
agg=defaultdict(lambda:[0,0,0.,0.]); seen=set(); rows=0; over100=0
for f in sorted((ROOT/'data').glob('products_*.csv.gz')):
    with gzip.open(f,'rt',encoding='utf-8',newline='') as h:
        for r in csv.DictReader(h):
            assert r['url'] not in seen; seen.add(r['url']); rows+=1
            assert int(r['source_row'])==rows+9
            rev=float(r['revenue']);drr=float(r['drr']);over100+=drr>1
            a=agg[tuple(r[c] for c in ['cat1','cat2','cat3','cat4'])]
            a[0]+=1;a[1]+=rev>1000000;a[2]+=rev;a[3]+=rev*drr
with (ROOT/'outputs/categories.csv').open(encoding='utf-8-sig',newline='') as h:
    result=list(csv.DictReader(h))
assert len(agg)==len(result)==4588 and rows==300000 and over100==37
selected=[]
for r in result:
    a=agg[tuple(r[c] for c in ['cat1','cat2','cat3','cat4'])]
    assert a[0]==int(r['n']) and a[1]==int(r['million_n'])
    assert math.isclose(a[2],float(r['revenue']),rel_tol=1e-12,abs_tol=.01)
    assert math.isclose(a[3]/a[2],float(r['weighted_drr']),rel_tol=1e-10,abs_tol=1e-12)
    assert math.isclose(a[1]/a[0],float(r['million_share']),abs_tol=1e-12)
    if a[0]>=30 and a[3]/a[2]<=.03 and a[1]/a[0]>=.1:selected.append(r['cat4'])
assert len(selected)==10
assert sum(a[1] for a in agg.values())==1476
book=ROOT/'outputs/Ozon_анализ.xlsx'
wb=openpyxl.load_workbook(book,read_only=True,data_only=True)
ws=wb['Источник']
assert ws['B20'].value==300000 and ws['B21'].value==1476
assert math.isclose(ws['B22'].value,sum(a[2] for a in agg.values()),abs_tol=.02)
wc=wb['Категории']
cached=list(wc.iter_rows(min_row=6,values_only=True))
assert len(cached)==4588
for a,r in zip(cached,result):
    assert math.isclose(a[13],float(r['weighted_drr']),abs_tol=1e-12)
    assert math.isclose(a[14],float(r['million_share']),abs_tol=1e-12)
    assert a[15]==('Да' if r['selected']=='True' else 'Нет')
report={'independent_streaming_controls':'passed','xlsx_cached_values':'passed','rows':rows,'categories':len(agg),'million_products':1476,'selected_categories':len(selected),'drr_over_100pct_retained':over100,'native_excel_tested':False}
(ROOT/'outputs/validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
